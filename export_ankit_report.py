import os, sys, django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

import zoneinfo
from datetime import datetime, time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.contrib.auth import get_user_model
from leads.models import Lead
from followups.models import FollowUp, Activity
from audit.models import AuditLog

User = get_user_model()
ankit = User.objects.filter(username='ankit_P').first()

if not ankit:
    print("User ankit_P not found!")
    sys.exit(1)

ist = zoneinfo.ZoneInfo('Asia/Kolkata')
today_ist = datetime.now(ist).date()
start_ist = datetime.combine(today_ist, time.min).replace(tzinfo=ist)
end_ist = datetime.combine(today_ist, time.max).replace(tzinfo=ist)

acts_today = Activity.objects.filter(created_by=ankit, created_at__gte=start_ist, created_at__lte=end_ist)
fus_created_today = FollowUp.objects.filter(created_by=ankit, created_at__gte=start_ist, created_at__lte=end_ist)
fus_updated_today = FollowUp.objects.filter(created_by=ankit, updated_at__gte=start_ist, updated_at__lte=end_ist)
audits_today = AuditLog.objects.filter(user=ankit, created_at__gte=start_ist, created_at__lte=end_ist)

lead_ids = set(acts_today.values_list('lead_id', flat=True))
lead_ids.update(fus_created_today.values_list('lead_id', flat=True))
lead_ids.update(fus_updated_today.values_list('lead_id', flat=True))

for a in audits_today:
    if a.model_name == 'Lead' and a.object_id:
        try:
            lead_ids.add(int(a.object_id))
        except:
            pass
    elif a.model_name == 'FollowUp' and a.object_id:
        fu = FollowUp.objects.filter(id=a.object_id).first()
        if fu:
            lead_ids.add(fu.lead_id)

leads = Lead.objects.filter(id__in=lead_ids).select_related('course', 'stage', 'lead_source', 'assigned_to').prefetch_related('followups', 'activities', 'lead_notes').order_by('id')

print(f"Total leads fetched: {len(leads)}")

wb = openpyxl.Workbook()

# Sheet 1: Leads Overview & Details
ws1 = wb.active
ws1.title = 'Leads Overview & Details'

header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_center = Alignment(horizontal='center', vertical='center')

headers_ws1 = [
    'S.No.', 'Lead ID', 'Lead Code', 'Student Name', 'Mobile Number', 'Email',
    'City', 'Course / Program', 'Stage', 'Temperature', 'Deal Status', 'Admission Status',
    'Inquiry Date', 'Next Followup Date', 'Next Followup Time', 'Total Followups Count',
    'Today Actions by Ankit', 'Latest Remark / Note', 'All Followups Summary (Chronological)'
]

ws1.append(headers_ws1)

for col_num in range(1, len(headers_ws1) + 1):
    cell = ws1.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = align_center

row_idx = 2
for idx, lead in enumerate(leads, start=1):
    actions = []
    lead_acts = [a for a in acts_today if a.lead_id == lead.id]
    lead_fus_c = [f for f in fus_created_today if f.lead_id == lead.id]
    lead_fus_u = [f for f in fus_updated_today if f.lead_id == lead.id]
    
    if lead_fus_c:
        actions.append(f'{len(lead_fus_c)} Follow-up(s) Added/Scheduled')
    if lead_fus_u:
        actions.append(f'{len(lead_fus_u)} Follow-up(s) Updated/Rescheduled')
    for la in lead_acts:
        if la.activity_type == 'STAGE_CHANGE':
            actions.append(f'Stage Changed: {la.description}')
        elif la.activity_type == 'TEMPERATURE_CHANGE':
            actions.append(f'Temperature Changed: {la.description}')
        elif la.activity_type == 'NOTE':
            actions.append(f'Note: {la.description}')
    
    actions_str = ' | '.join(list(dict.fromkeys(actions))) if actions else 'Lead Profile / Followup Updated'
    
    all_fu_list = lead.followups.all().order_by('followup_date', 'id')
    fu_summary_parts = []
    for f in all_fu_list:
        f_date = str(f.followup_date) if f.followup_date else '-'
        f_status = f.get_followup_status_display() if hasattr(f, 'get_followup_status_display') else f.followup_status
        f_mode = f.get_followup_mode_display() if hasattr(f, 'get_followup_mode_display') else f.followup_mode
        f_comment = f.comment.strip() if f.comment else 'No remark'
        fu_summary_parts.append(f'[{f_date} | {f_mode} | {f_status}]: {f_comment}')
    
    fu_summary_str = '\n'.join(fu_summary_parts) if fu_summary_parts else 'No follow-ups recorded'
    latest_remark = lead.notes if lead.notes else (all_fu_list.last().comment if all_fu_list.exists() and all_fu_list.last().comment else '-')

    row_data = [
        idx,
        lead.id,
        lead.lead_code or '-',
        lead.name or '-',
        lead.mobile or '-',
        lead.email or '-',
        lead.city or '-',
        lead.course.name if lead.course else '-',
        lead.stage.name if lead.stage else '-',
        lead.get_temperature_display() if hasattr(lead, 'get_temperature_display') else lead.temperature,
        lead.get_deal_status_display() if hasattr(lead, 'get_deal_status_display') else lead.deal_status,
        lead.get_admission_status_display() if hasattr(lead, 'get_admission_status_display') else lead.admission_status,
        str(lead.inquiry_date) if lead.inquiry_date else '-',
        str(lead.next_followup_date) if lead.next_followup_date else '-',
        str(lead.next_followup_time) if lead.next_followup_time else '-',
        lead.followups.count(),
        actions_str,
        latest_remark,
        fu_summary_str
    ]
    ws1.append(row_data)
    for col_num in range(1, len(row_data) + 1):
        c = ws1.cell(row=row_idx, column=col_num)
        c.border = border_thin
        c.alignment = align_center if col_num in [1, 2, 9, 10, 11, 12, 13, 14, 15, 16] else align_left
    row_idx += 1

# Sheet 2: All Follow-ups Detailed Breakdown
ws2 = wb.create_sheet(title='All Follow-ups Detailed')
header_fill_2 = PatternFill(start_color='203764', end_color='203764', fill_type='solid')
headers_ws2 = [
    'S.No.', 'FollowUp ID', 'Lead ID', 'Lead Code', 'Student Name', 'Mobile Number',
    'Course', 'Lead Stage', 'Temperature', 'Deal Status',
    'Follow-up Date', 'Follow-up Time', 'Mode', 'Follow-up Status',
    'Follow-up Remarks / Comment', 'Next Follow-up Date Scheduled',
    'Created By', 'Created At (IST)', 'Last Updated At (IST)', 'Touched Today by Ankit?'
]
ws2.append(headers_ws2)
for col_num in range(1, len(headers_ws2) + 1):
    cell = ws2.cell(row=1, column=col_num)
    cell.fill = header_fill_2
    cell.font = header_font
    cell.alignment = align_center

fu_row_idx = 2
fu_sno = 1

for lead in leads:
    for fu in lead.followups.all().order_by('followup_date', 'id'):
        c_at_ist = fu.created_at.astimezone(ist).strftime('%Y-%m-%d %I:%M %p') if fu.created_at else '-'
        u_at_ist = fu.updated_at.astimezone(ist).strftime('%Y-%m-%d %I:%M %p') if fu.updated_at else '-'
        
        is_today = 'Yes' if (fu in fus_created_today or fu in fus_updated_today) else 'No'
        
        fu_data = [
            fu_sno,
            fu.id,
            lead.id,
            lead.lead_code or '-',
            lead.name or '-',
            lead.mobile or '-',
            lead.course.name if lead.course else '-',
            lead.stage.name if lead.stage else '-',
            lead.get_temperature_display() if hasattr(lead, 'get_temperature_display') else lead.temperature,
            lead.get_deal_status_display() if hasattr(lead, 'get_deal_status_display') else lead.deal_status,
            str(fu.followup_date) if fu.followup_date else '-',
            str(fu.followup_time) if fu.followup_time else '-',
            fu.get_followup_mode_display() if hasattr(fu, 'get_followup_mode_display') else fu.followup_mode,
            fu.get_followup_status_display() if hasattr(fu, 'get_followup_status_display') else fu.followup_status,
            fu.comment or '-',
            str(fu.next_followup_date) if fu.next_followup_date else '-',
            fu.created_by.get_full_name() if fu.created_by else 'System / Unassigned',
            c_at_ist,
            u_at_ist,
            is_today
        ]
        ws2.append(fu_data)
        for col_num in range(1, len(fu_data) + 1):
            c = ws2.cell(row=fu_row_idx, column=col_num)
            c.border = border_thin
            c.alignment = align_center if col_num in [1, 2, 3, 4, 8, 9, 10, 11, 12, 13, 14, 16, 18, 19, 20] else align_left
        fu_row_idx += 1
        fu_sno += 1

# Sheet 3: Today Activity Log of Ankit
ws3 = wb.create_sheet(title='Ankit Today Activity Timeline')
header_fill_3 = PatternFill(start_color='333f48', end_color='333f48', fill_type='solid')
headers_ws3 = [
    'S.No.', 'Activity Time (IST)', 'Lead ID', 'Lead Code', 'Student Name', 'Mobile Number',
    'Activity Type', 'Description / Action Done'
]
ws3.append(headers_ws3)
for col_num in range(1, len(headers_ws3) + 1):
    cell = ws3.cell(row=1, column=col_num)
    cell.fill = header_fill_3
    cell.font = header_font
    cell.alignment = align_center

act_row_idx = 2
for idx, act in enumerate(acts_today.select_related('lead').order_by('created_at'), start=1):
    t_ist = act.created_at.astimezone(ist).strftime('%Y-%m-%d %I:%M:%S %p') if act.created_at else '-'
    act_data = [
        idx,
        t_ist,
        act.lead.id if act.lead else '-',
        act.lead.lead_code if act.lead else '-',
        act.lead.name if act.lead else '-',
        act.lead.mobile if act.lead else '-',
        act.get_activity_type_display() if hasattr(act, 'get_activity_type_display') else act.activity_type,
        act.description
    ]
    ws3.append(act_data)
    for col_num in range(1, len(act_data) + 1):
        c = ws3.cell(row=act_row_idx, column=col_num)
        c.border = border_thin
        c.alignment = align_center if col_num in [1, 2, 3, 4, 7] else align_left
    act_row_idx += 1

# Auto-adjust column widths
for ws in [ws1, ws2, ws3]:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            lines = val_str.split('\n')
            max_line_len = max(len(l) for l in lines) if lines else 0
            if max_line_len > max_len:
                max_len = max_line_len
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

filename = f'Ankit_Pardhi_Activity_Report_Today_{today_ist}.xlsx'
output_path = os.path.abspath(filename)
wb.save(output_path)
print(f"SUCCESS: File saved to {output_path}")
