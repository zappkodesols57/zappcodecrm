import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# -------------------------------------------------------------
# Color Palette (Healthcare Professional)
# -------------------------------------------------------------
NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A (Deep Slate/Navy)
NAVY_CARD = RGBColor(30, 41, 59)      # #1E293B (Card in dark slide)
SLATE_BG = RGBColor(248, 250, 252)    # #F8FAFC (Light slide background)
CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
BORDER_LIGHT = RGBColor(226, 232, 240)# #E2E8F0
PRIMARY_BLUE = RGBColor(37, 99, 235)  # #2563EB
CYAN_ACCENT = RGBColor(14, 165, 233)  # #0EA5E9
TEXT_DARK = RGBColor(15, 23, 42)      # #0F172A
TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
TEXT_LIGHT = RGBColor(241, 245, 249)  # #F1F5F9
ACCENT_RED = RGBColor(239, 68, 68)    # #EF4444 (Problem callouts)
ACCENT_GREEN = RGBColor(16, 185, 129) # #10B981 (Success/Impact)
ACCENT_AMBER = RGBColor(245, 158, 11) # #F59E0B

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

def set_slide_bg(slide, color):
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = color
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, badge_text, title_text, subtitle_text, is_dark=False):
    # Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(3.2), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 58, 138) if is_dark else RGBColor(238, 242, 255)
    badge.line.fill.background()
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = badge_text.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.name = "Calibri"
    p_b.font.color.rgb = CYAN_ACCENT if is_dark else PRIMARY_BLUE

    # Title & Subtitle Box
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(11.7), Inches(1.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_t = tf.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.name = "Calibri"
    p_t.font.color.rgb = TEXT_LIGHT if is_dark else TEXT_DARK

    p_s = tf.add_paragraph()
    p_s.text = subtitle_text
    p_s.font.size = Pt(13)
    p_s.font.name = "Calibri"
    p_s.font.color.rgb = CYAN_ACCENT if is_dark else TEXT_MUTED
    p_s.space_before = Pt(4)


# ==============================================================================
# SLIDE 1: Title Slide (Dark Theme)
# ==============================================================================
s1 = prs.slides.add_slide(blank_slide_layout)
set_slide_bg(s1, NAVY_DARK)

# Decorative accent bar
bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(1.2), Inches(0.1))
bar.fill.solid()
bar.fill.fore_color.rgb = CYAN_ACCENT
bar.line.fill.background()

# Title text frame
tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.85), Inches(11.5), Inches(3.2))
tf1 = tb1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

p1 = tf1.paragraphs[0]
p1.text = "Nelson Hospital"
p1.font.size = Pt(40)
p1.font.bold = True
p1.font.name = "Calibri"
p1.font.color.rgb = TEXT_LIGHT

p2 = tf1.add_paragraph()
p2.text = "Healthcare CRM Transformation"
p2.font.size = Pt(36)
p2.font.bold = True
p2.font.name = "Calibri"
p2.font.color.rgb = CYAN_ACCENT
p2.space_before = Pt(6)

p3 = tf1.add_paragraph()
p3.text = "Eliminating Lead Leakage • Automating Follow-Ups • Delivering Predictable Patient Conversions"
p3.font.size = Pt(16)
p3.font.name = "Calibri"
p3.font.color.rgb = RGBColor(148, 163, 184)
p3.space_before = Pt(16)

# Bottom Info Card
b_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.2))
b_card.fill.solid()
b_card.fill.fore_color.rgb = NAVY_CARD
b_card.line.color.rgb = RGBColor(51, 65, 85)
b_card.line.width = Pt(1)

tf_b = b_card.text_frame
tf_b.word_wrap = True
tf_b.margin_left = Inches(0.3)
tf_b.margin_top = Inches(0.2)

p_b1 = tf_b.paragraphs[0]
p_b1.text = "Strategy & Implementation Proposal"
p_b1.font.bold = True
p_b1.font.size = Pt(13)
p_b1.font.color.rgb = CYAN_ACCENT

p_b2 = tf_b.add_paragraph()
p_b2.text = "Prepared by: Zappcode Solutions  |  Target: Nelson Hospital Administration, Doctors & Lead Attendants"
p_b2.font.size = Pt(12)
p_b2.font.color.rgb = RGBColor(203, 213, 225)
p_b2.space_before = Pt(4)


# ==============================================================================
# SLIDE 2: Current Challenges (Problem Statement)
# ==============================================================================
s2 = prs.slides.add_slide(blank_slide_layout)
set_slide_bg(s2, SLATE_BG)
add_header(s2, "Current Operational Bottlenecks", "Where Hospital Revenue & Patient Inquiries Are Leaking", "The core issues observed in traditional manual and spreadsheet-based patient tracking")

cards_data_s2 = [
    ("1. Scattered & Missing Patient Data", 
     "Patient leads reside across detached Excel workbooks, paper registers, and staff phone logs.\n• Critical contact details and symptoms get lost between shifts.\n• No central record exists for a patient's historical inquiry history.",
     ACCENT_RED),
    ("2. Missed & Forgotten Follow-Ups",
     "Telecallers manually manage hundreds of follow-up dates without scheduled reminders.\n• High-intent inquiries go cold without timely second and third touches.\n• Inability to track whether an inquiry was actually called or abandoned.",
     ACCENT_AMBER),
    ("3. Month-End Results Disconnect",
     "Significant marketing spend on campaigns fails to translate to expected OPD footfalls.\n• At month-end, reported inquiry volume does not match booked doctor consultations.\n• Lack of visibility into where prospects dropped off in the funnel.",
     PRIMARY_BLUE),
    ("4. Source & Attribution Blindness",
     "Management cannot verify which channel drove which patient admission.\n• Unable to differentiate organic walk-ins from paid Meta Ads or Health Camps.\n• Marketing budget gets allocated blindly without knowing real departmental ROI.",
     RGBColor(124, 58, 237))
]

for idx, (title, desc, accent) in enumerate(cards_data_s2):
    row = idx // 2
    col = idx % 2
    x = Inches(0.8 + col * 5.95)
    y = Inches(2.2 + row * 2.4)

    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.15))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_LIGHT
    card.line.width = Pt(1)

    # Accent line top of card
    top_line = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.2), y + Inches(0.2), Inches(0.5), Inches(0.06))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = accent
    top_line.line.fill.background()

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.35)

    p_t = tf.paragraphs[0]
    p_t.text = title
    p_t.font.bold = True
    p_t.font.size = Pt(14)
    p_t.font.color.rgb = TEXT_DARK

    p_d = tf.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(11)
    p_d.font.color.rgb = TEXT_MUTED
    p_d.space_before = Pt(6)


# ==============================================================================
# SLIDE 3: The Solution (Centralized Hospital CRM)
# ==============================================================================
s3 = prs.slides.add_slide(blank_slide_layout)
set_slide_bg(s3, SLATE_BG)
add_header(s3, "The Solution Architecture", "Zappcode Healthcare CRM: Built for Nelson Hospital", "A unified digital platform connecting marketing, front desk, telecallers, and doctors")

pillars = [
    ("Real-Time Automated Capture",
     "Direct API Webhooks",
     "Zero manual entry. Inquiries from Meta Ads, Google campaigns, website forms, and walk-in counters land instantly in the CRM database within seconds.\n\n• Instant lead ingestion\n• Duplicate mobile number detection\n• Notification alerts to counselors",
     CYAN_ACCENT),
    ("Departmental Segmentation",
     "Specialty-Based Routing",
     "Automated routing of patients directly into Nelson's specialized branches and wings.\n\n• Neuro Sciences\n• Gynaecology & Obstetrics\n• Nelson Luxe Care Suites\n• Health Camps & General OPD\n• Custom medical form fields",
     PRIMARY_BLUE),
    ("End-to-End Patient Journey",
     "Full Lifecycle Visibility",
     "Track every patient from discovery to full recovery across dedicated milestone stages.\n\n• New Lead Received\n• Telecaller Follow-up\n• Doctor OPD Appointment Booked\n• Completed Consultation\n• In-Patient Admission / Procedure",
     ACCENT_GREEN)
]

for idx, (title, tag, body, accent) in enumerate(pillars):
    x = Inches(0.8 + idx * 4.0)
    y = Inches(2.2)

    card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.75), Inches(4.7))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_LIGHT
    card.line.width = Pt(1)

    # Accent Pill Tag
    tag_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.3), y + Inches(0.3), Inches(2.4), Inches(0.32))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    tag_box.line.color.rgb = accent
    tag_box.line.width = Pt(1)
    p_tag = tag_box.text_frame.paragraphs[0]
    p_tag.text = tag.upper()
    p_tag.font.size = Pt(9)
    p_tag.font.bold = True
    p_tag.font.color.rgb = accent

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.75)

    p_t = tf.paragraphs[0]
    p_t.text = title
    p_t.font.bold = True
    p_t.font.size = Pt(15)
    p_t.font.color.rgb = TEXT_DARK

    p_b = tf.add_paragraph()
    p_b.text = body
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = TEXT_MUTED
    p_b.space_before = Pt(10)


# ==============================================================================
# SLIDE 4: Instant Patient Engagement (WhatsApp & Calling)
# ==============================================================================
s4 = prs.slides.add_slide(blank_slide_layout)
set_slide_bg(s4, SLATE_BG)
add_header(s4, "Instant Patient Engagement", "1-Click Direct WhatsApp & Automated Follow-Up Queue", "Empowering telecallers to engage warm inquiries immediately without manual dialing delays")

# Left Box: Direct WhatsApp & Communication
card_left = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.75), Inches(4.7))
card_left.fill.solid()
card_left.fill.fore_color.rgb = CARD_BG
card_left.line.color.rgb = BORDER_LIGHT
card_left.line.width = Pt(1)

tf_l = card_left.text_frame
tf_l.word_wrap = True
tf_l.margin_left = Inches(0.4)
tf_l.margin_right = Inches(0.4)
tf_l.margin_top = Inches(0.4)

p_lt = tf_l.paragraphs[0]
p_lt.text = "💬 1-Click WhatsApp & Direct Calling"
p_lt.font.bold = True
p_lt.font.size = Pt(16)
p_lt.font.color.rgb = PRIMARY_BLUE

p_lb = tf_l.add_paragraph()
p_lb.text = (
    "• No Number Saving Required:\n"
    "  Staff clicks one button on the patient row to immediately trigger official WhatsApp chat or phone dialer.\n\n"
    "• Pre-Approved Hospital Templates:\n"
    "  Send instant welcome messages, hospital location maps, doctor profiles, and appointment confirmations.\n\n"
    "• Interaction Timeline:\n"
    "  Every communication attempt, remark, and feedback is recorded on the patient's card with timestamps."
)
p_lb.font.size = Pt(12)
p_lb.font.color.rgb = TEXT_MUTED
p_lb.space_before = Pt(12)

# Right Box: Dynamic Calling Queue
card_right = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.78), Inches(2.2), Inches(5.75), Inches(4.7))
card_right.fill.solid()
card_right.fill.fore_color.rgb = CARD_BG
card_right.line.color.rgb = BORDER_LIGHT
card_right.line.width = Pt(1)

tf_r = card_right.text_frame
tf_r.word_wrap = True
tf_r.margin_left = Inches(0.4)
tf_r.margin_right = Inches(0.4)
tf_r.margin_top = Inches(0.4)

p_rt = tf_r.paragraphs[0]
p_rt.text = "⏰ Zero-Leakage Follow-Up Queue"
p_rt.font.bold = True
p_rt.font.size = Pt(16)
p_rt.font.color.rgb = ACCENT_GREEN

p_rb = tf_r.add_paragraph()
p_rb.text = (
    "• Daily Pending Calling Dashboard:\n"
    "  Telecallers start their day with a clear, filtered queue of overdue calls and today's scheduled follow-ups.\n\n"
    "• Smart Remark Logging:\n"
    "  Enforces structured call remarks (Remark 1, Remark 2, Doctor preference) to prevent abandoned leads.\n\n"
    "• High-Speed Patient Response:\n"
    "  Responding within 10 minutes increases OPD appointment booking rates by up to 300%."
)
p_rb.font.size = Pt(12)
p_rb.font.color.rgb = TEXT_MUTED
p_rb.space_before = Pt(12)


# ==============================================================================
# SLIDE 5: Campaign Intelligence & Marketing Optimization
# ==============================================================================
s5 = prs.slides.add_slide(blank_slide_layout)
set_slide_bg(s5, SLATE_BG)
add_header(s5, "Strategic Marketing Insights", "Turning Ad Spend into Measurable Hospital Admissions", "Gain complete clarity on which campaigns produce actual patients and where to invest next")

metrics_data = [
    ("Granular Campaign Tracking",
     "Real-Time Attribution",
     "Track every single patient back to the exact Meta Ad campaign, ad set, creative, or offline health camp banner.\n\n• Know cost-per-lead by specialty\n• Eliminate guesswork in ad agency reports\n• Validate real patient inquiries vs bot clicks",
     CYAN_ACCENT),
    ("Specialty Demand Analysis",
     "Departmental Insights",
     "Understand which medical sectors are generating high patient interest across regions.\n\n• High conversion: Neuro & Gynae\n• High volume: General Health Camps\n• Premium inquiries: Luxe Care Suites\n• Direct marketing efforts to high-margin treatments",
     PRIMARY_BLUE),
    ("Predictable Month-End Forecasting",
     "Reliable Hospital Growth",
     "Match marketing budget directly against hospital footfall and revenue conversion.\n\n• Track conversion ratios: Inquiries ➔ Consultations\n• Clear insight on uncontacted leads\n• Strategic planning for future medical campaigns",
     RGBColor(147, 51, 234))
]

for idx, (title, sub, body, accent) in enumerate(metrics_data):
    x = Inches(0.8 + idx * 4.0)
    y = Inches(2.2)

    card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.75), Inches(4.7))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_LIGHT
    card.line.width = Pt(1)

    bar = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.3), y + Inches(0.3), Inches(0.8), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.55)

    p_t = tf.paragraphs[0]
    p_t.text = title
    p_t.font.bold = True
    p_t.font.size = Pt(15)
    p_t.font.color.rgb = TEXT_DARK

    p_s = tf.add_paragraph()
    p_s.text = sub.upper()
    p_s.font.size = Pt(9)
    p_s.font.bold = True
    p_s.font.color.rgb = accent
    p_s.space_before = Pt(3)

    p_b = tf.add_paragraph()
    p_b.text = body
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = TEXT_MUTED
    p_b.space_before = Pt(10)


# ==============================================================================
# SLIDE 6: Role-Based Workflow for Nelson Hospital
# ==============================================================================
s6 = prs.slides.add_slide(blank_slide_layout)
set_slide_bg(s6, SLATE_BG)
add_header(s6, "Collaborative Hospital Ecosystem", "Tailored Interfaces Designed for Every Role", "Ensuring every hospital department has exactly the tools and visibility they need")

roles_data = [
    ("Executive Leadership & Admin",
     "Management / Directors",
     "• Hospital-wide macro performance dashboard\n"
     "• Real-time lead volume, conversion %, and revenue tracking\n"
     "• Campaign ROI & agency accountability\n"
     "• Master settings (Doctors, Departments, Branches, Users)",
     NAVY_DARK),
    ("Telecallers & Front Desk",
     "Lead Attendants / Reception",
     "• Clean personal calling queue with priority sorting\n"
     "• 1-Click WhatsApp greeting & patient outreach\n"
     "• Structured call remarks & deal status tracking\n"
     "• Direct OPD appointment booking into doctor calendars",
     PRIMARY_BLUE),
    ("Doctors & Consultants",
     "Specialists / Medical Staff",
     "• Daily OPD consultation schedule & patient lists\n"
     "• Real-time visibility into booked appointment slots\n"
     "• Easy leave scheduling & available hours management\n"
     "• Patient interaction and referral notes",
     CYAN_ACCENT)
]

for idx, (title, target, items, color) in enumerate(roles_data):
    y = Inches(2.2 + idx * 1.55)

    card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.73), Inches(1.38))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_LIGHT
    card.line.width = Pt(1)

    # Left indicator bar
    side_bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y, Inches(0.12), Inches(1.38))
    side_bar.fill.solid()
    side_bar.fill.fore_color.rgb = color
    side_bar.line.fill.background()

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.35)
    tf.margin_top = Inches(0.15)
    tf.margin_right = Inches(0.3)

    p_t = tf.paragraphs[0]
    p_t.text = f"{title}  ({target})"
    p_t.font.bold = True
    p_t.font.size = Pt(13)
    p_t.font.color.rgb = color

    p_i = tf.add_paragraph()
    p_i.text = items
    p_i.font.size = Pt(10.5)
    p_i.font.color.rgb = TEXT_MUTED
    p_i.space_before = Pt(3)


# ==============================================================================
# SLIDE 7: Conclusion & Summary (Dark Theme)
# ==============================================================================
s7 = prs.slides.add_slide(blank_slide_layout)
set_slide_bg(s7, NAVY_DARK)
add_header(s7, "The Measurable Impact", "Summary: The Transformation for Nelson Hospital", "Moving from manual friction to scalable, modern healthcare operations", is_dark=True)

kpi_data = [
    ("100%", "Captured Leads", "Zero inquiries lost or missing from Meta Ads & walk-ins", ACCENT_GREEN),
    ("3X", "Faster Response", "Immediate WhatsApp outreach while patient intent is hot", CYAN_ACCENT),
    ("0%", "Forgotten Calls", "Automated daily follow-up queues for every telecaller", PRIMARY_BLUE),
    ("Full", "ROI Transparency", "Clear attribution of marketing budget to admitted patients", RGBColor(168, 85, 247))
]

for idx, (num, label, desc, accent) in enumerate(kpi_data):
    x = Inches(0.8 + idx * 3.0)
    y = Inches(2.2)

    card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.78), Inches(2.6))
    card.fill.solid()
    card.fill.fore_color.rgb = NAVY_CARD
    card.line.color.rgb = RGBColor(51, 65, 85)
    card.line.width = Pt(1)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)

    p_num = tf.paragraphs[0]
    p_num.text = num
    p_num.font.bold = True
    p_num.font.size = Pt(32)
    p_num.font.color.rgb = accent

    p_lbl = tf.add_paragraph()
    p_lbl.text = label
    p_lbl.font.bold = True
    p_lbl.font.size = Pt(13)
    p_lbl.font.color.rgb = TEXT_LIGHT
    p_lbl.space_before = Pt(4)

    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.size = Pt(10.5)
    p_desc.font.color.rgb = RGBColor(148, 163, 184)
    p_desc.space_before = Pt(6)

# Closing Summary Banner
banner = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.1), Inches(11.73), Inches(1.6))
banner.fill.solid()
banner.fill.fore_color.rgb = NAVY_CARD
banner.line.color.rgb = CYAN_ACCENT
banner.line.width = Pt(1.5)

tf_ban = banner.text_frame
tf_ban.word_wrap = True
tf_ban.margin_left = Inches(0.4)
tf_ban.margin_top = Inches(0.25)
tf_ban.margin_right = Inches(0.4)

p_b1 = tf_ban.paragraphs[0]
p_b1.text = "Empowering Nelson Hospital for Scalable Growth"
p_b1.font.bold = True
p_b1.font.size = Pt(15)
p_b1.font.color.rgb = CYAN_ACCENT

p_b2 = tf_ban.add_paragraph()
p_b2.text = (
    "With Zappcode CRM, Nelson Hospital eliminates lead leakage, ensures every warm patient is contacted "
    "systematically via WhatsApp and phone, and gives management the exact data needed to expand key departments "
    "(Neuro, Gynaecology, Luxe Care). Let's build a smarter, patient-first healthcare workflow together!"
)
p_b2.font.size = Pt(11.5)
p_b2.font.color.rgb = RGBColor(226, 232, 240)
p_b2.space_before = Pt(6)


# -------------------------------------------------------------
# Save Presentation
# -------------------------------------------------------------
output_path = r"F:\untitled folder\zappcodecrm\Nelson_Hospital_CRM_Presentation.pptx"
prs.save(output_path)
print(f"SUCCESS: Presentation created at {output_path}")

